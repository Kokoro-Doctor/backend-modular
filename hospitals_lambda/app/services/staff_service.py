"""
Staff service - handles adding patients and doctors on behalf of hospital staff.

Creates user/doctor profiles, auth records, and hospital linkages.
Follows the same patterns as auth_lambda's user_service and doctor_service.
"""
import io
import re
import uuid
from datetime import datetime, timezone, timedelta, time
from typing import Any, Dict, Optional

from fastapi import HTTPException, UploadFile
from boto3.dynamodb.conditions import Key

from app.config import (
    USERS_TABLE,
    DOCTORS_TABLE,
    AUTH_TABLE,
    DOCTOR_AVAILABILITY_TABLE,
    SMS_COUNTRY_CODE,
)
from app.services.user_doctor_relation_service import (
    create_relation,
    deactivate_relation,
    get_active_hospital_assigned_doctor,
    RelationType,
    LinkedBy,
)
from app.logger import get_logger

logger = get_logger(__name__)


def _phone_tail(phone: str) -> str:
    """Last 4 digits for logs (avoid logging full numbers)."""
    digits = re.sub(r"\D", "", phone or "")
    if len(digits) >= 4:
        return f"***{digits[-4:]}"
    return "***"


# ---------------------------------------------------------------------------
# Phone normalization (mirrors auth_lambda/app/utils/db_utils.py)
# ---------------------------------------------------------------------------

def normalize_phone_number(phone: str) -> str:
    """Normalize phone number to E.164 format."""
    if not phone:
        return ""
    trimmed = phone.strip()
    if not trimmed:
        return ""
    digits_only = re.sub(r"\D", "", trimmed)
    if not digits_only:
        return ""

    if trimmed.startswith("+"):
        normalized = "+" + digits_only
        if len(normalized) < 8 or len(normalized) > 18:
            return ""
        if digits_only.startswith("91"):
            if len(digits_only) == 12:
                return normalized
            if len(digits_only) == 11:
                return f"+91{digits_only[-10:]}"
            if len(digits_only) == 10:
                return f"+91{digits_only}"
        return normalized

    if len(digits_only) == 10:
        return f"{SMS_COUNTRY_CODE}{digits_only}"
    if len(digits_only) == 12 and digits_only.startswith("91"):
        return f"+{digits_only}"
    if len(digits_only) == 11 and digits_only.startswith("91"):
        return f"{SMS_COUNTRY_CODE}{digits_only[-10:]}"
    if len(digits_only) > 12:
        last_12 = digits_only[-12:]
        if last_12.startswith("91"):
            return f"+{last_12}"
        return f"{SMS_COUNTRY_CODE}{digits_only[-10:]}"
    if len(digits_only) < 10:
        return ""

    return f"{SMS_COUNTRY_CODE}{digits_only[-10:]}"


# ---------------------------------------------------------------------------
# ID generation (mirrors auth_lambda/app/utils/db_utils.py)
# ---------------------------------------------------------------------------

def _prefixed_id(prefix: str) -> str:
    return f"{prefix}_{uuid.uuid4()}"


def generate_user_id() -> str:
    return _prefixed_id("usr")


def generate_doctor_id() -> str:
    return _prefixed_id("dr")


# ---------------------------------------------------------------------------
# Lookup helpers
# ---------------------------------------------------------------------------

def _get_user_by_phone(phone: str) -> Optional[dict]:
    normalized = normalize_phone_number(phone)
    if not normalized:
        return None
    try:
        resp = USERS_TABLE.query(
            IndexName="phone-index",
            KeyConditionExpression=Key("phoneNumber").eq(normalized),
        )
        items = resp.get("Items", [])
        return items[0] if items else None
    except Exception as e:
        logger.error(f"[staff] get_user_by_phone error: {e}")
        return None


def _get_user_by_id(user_id: str) -> Optional[dict]:
    if not user_id or not str(user_id).strip():
        return None
    try:
        resp = USERS_TABLE.get_item(Key={"user_id": str(user_id).strip()})
        return resp.get("Item")
    except Exception as e:
        logger.error(f"[staff] get_user_by_id error user_id={user_id!r}: {e}")
        return None


def _get_doctor_by_phone(phone: str) -> Optional[dict]:
    normalized = normalize_phone_number(phone)
    if not normalized:
        return None
    try:
        resp = DOCTORS_TABLE.query(
            IndexName="phone-index",
            KeyConditionExpression=Key("phoneNumber").eq(normalized),
        )
        items = resp.get("Items", [])
        return items[0] if items else None
    except Exception as e:
        logger.error(f"[staff] get_doctor_by_phone error: {e}")
        return None


def get_doctor_for_hospital_staff_patient_flow(doctor_id: str) -> dict:
    """
    Load doctor by id for staff add-patient / import-patient flows.
    404 if missing; 400 if doctor has no hospital affiliation.
    """
    if not doctor_id or not str(doctor_id).strip():
        raise HTTPException(status_code=400, detail="doctor_id is required")
    did = str(doctor_id).strip()
    try:
        resp = DOCTORS_TABLE.get_item(Key={"doctor_id": did})
    except Exception as e:
        logger.error(f"[staff] get_doctor_for_hospital_staff_patient_flow get_item failed doctor_id={did!r} error={e!r}")
        raise HTTPException(status_code=500, detail="Failed to load doctor")
    doc = resp.get("Item")
    if not doc:
        raise HTTPException(status_code=404, detail="Doctor not found")
    hid = doc.get("hospital_id")
    if not hid:
        raise HTTPException(
            status_code=400,
            detail="Doctor is not affiliated with a hospital; hospital staff cannot add patients for this doctor",
        )
    return doc


# ---------------------------------------------------------------------------
# Auth record helpers (mirrors auth_lambda ensure_auth_record pattern)
# ---------------------------------------------------------------------------

def _ensure_auth_record(
    phone: str,
    role: str,
    user_id: Optional[str] = None,
    doctor_id: Optional[str] = None,
    email: Optional[str] = None,
) -> dict:
    """Create or update an AuthTable shell record so the person can log in via OTP."""
    normalized = normalize_phone_number(phone)
    if not normalized:
        logger.warning(f"[staff] _ensure_auth_record invalid phone raw={phone!r}")
        raise ValueError(f"Invalid phone number: {phone}")

    tail = _phone_tail(normalized)
    try:
        resp = AUTH_TABLE.get_item(Key={"phoneNumber": normalized})
        existing = resp.get("Item")
    except Exception as e:
        logger.error(f"[staff] AuthTable get_item failed phone={tail} error={e!r}")
        existing = None

    now = datetime.now(timezone.utc).isoformat()

    if existing:
        updates = {"updated_at": now}
        if role and not existing.get("role"):
            updates["role"] = role
        if user_id and not existing.get("user_id"):
            updates["user_id"] = user_id
        if doctor_id and not existing.get("doctor_id"):
            updates["doctor_id"] = doctor_id
        if email and not existing.get("email"):
            updates["email"] = email.lower().strip()

        if len(updates) > 1:
            try:
                expr_parts, attr_vals, attr_names = [], {}, {}
                for idx, (field, value) in enumerate(updates.items()):
                    name_ph = f"#f{idx}"
                    val_ph = f":v{idx}"
                    expr_parts.append(f"{name_ph} = {val_ph}")
                    attr_names[name_ph] = field
                    attr_vals[val_ph] = value
                AUTH_TABLE.update_item(
                    Key={"phoneNumber": normalized},
                    UpdateExpression="SET " + ", ".join(expr_parts),
                    ExpressionAttributeNames=attr_names,
                    ExpressionAttributeValues=attr_vals,
                )
                logger.info(
                    f"[staff] AuthTable updated phone={tail} role={role} "
                    f"fields={list(updates.keys())}"
                )
            except Exception as e:
                logger.error(f"[staff] AuthTable update_item failed phone={tail} error={e!r}")
                raise
        return existing

    record = {
        "phoneNumber": normalized,
        "role": role,
        "is_verified": False,
        "email_verified": False,
        "phone_verified": False,
        "user_id": user_id,
        "doctor_id": doctor_id,
        "created_at": now,
        "updated_at": now,
        "last_login": None,
        "last_verified_at": None,
    }
    if email:
        record["email"] = email.lower().strip()

    try:
        AUTH_TABLE.put_item(Item=record)
        logger.info(
            f"[staff] AuthTable created phone={tail} role={role} "
            f"user_id={user_id!r} doctor_id={doctor_id!r}"
        )
    except Exception as e:
        logger.error(f"[staff] AuthTable put_item failed phone={tail} role={role} error={e!r}")
        raise
    return record


# ---------------------------------------------------------------------------
# Doctor availability slots (mirrors auth_lambda doctor_service)
# ---------------------------------------------------------------------------

def _get_expiry_timestamp(date_str: str) -> int:
    date_obj = datetime.strptime(date_str, "%Y-%m-%d")
    end_of_day = datetime.combine(date_obj.date(), time(23, 59, 59))
    return int(end_of_day.timestamp())


def _create_default_slots(doctor_id: str) -> None:
    """Create default 9 AM - 5 PM / 30-min slots for the next 7 days."""
    try:
        today = datetime.now(timezone.utc).date()
        dates = [today + timedelta(days=i) for i in range(7)]
        slot_times = [
            "09:00", "09:30", "10:00", "10:30", "11:00", "11:30",
            "12:00", "12:30", "13:00", "13:30", "14:00", "14:30",
            "15:00", "15:30", "16:00", "16:30", "17:00",
        ]
        created_at = datetime.utcnow().isoformat()
        for date_obj in dates:
            date_str = date_obj.strftime("%Y-%m-%d")
            expiry = _get_expiry_timestamp(date_str)
            for st in slot_times:
                DOCTOR_AVAILABILITY_TABLE.put_item(Item={
                    "PK": doctor_id,
                    "SK": f"{date_str}#{st}",
                    "available": True,
                    "created_at": created_at,
                    "expiry_timestamp": expiry,
                })
        logger.info(f"[staff] Created default slots for doctor {doctor_id}")
    except Exception as e:
        logger.error(f"[staff] Failed to create default slots for {doctor_id}: {e}")


# ---------------------------------------------------------------------------
# Core: add patient
# ---------------------------------------------------------------------------


def _ensure_user_hospital_affiliation(
    user_id: str, hospital_id: str, hospital_name: str
) -> None:
    """Set Users.hospital_id / hospital_name from the hospital staff session."""
    try:
        USERS_TABLE.update_item(
            Key={"user_id": user_id},
            UpdateExpression="SET hospital_id = :hid, hospital_name = :hname",
            ExpressionAttributeValues={
                ":hid": hospital_id,
                ":hname": hospital_name,
            },
        )
        logger.info(
            f"[staff] set user hospital_affiliation user_id={user_id!r} "
            f"hospital_id={hospital_id!r}"
        )
    except Exception as e:
        logger.error(
            f"[staff] Users update_item hospital affiliation failed user_id={user_id!r} error={e!r}"
        )
        raise HTTPException(status_code=500, detail="Failed to set patient hospital affiliation")


def add_patient(
    phone: str,
    name: str,
    email: Optional[str] = None,
    *,
    hospital_id: str,
    hospital_name: str = "",
    doctor_id: Optional[str] = None,
    preloaded_doctor: Optional[dict] = None,
    age: Optional[int] = None,
    gender: Optional[str] = None,
    insurer: Optional[str] = None,
) -> dict:
    """
    With doctor_id: create/link a patient to that doctor. hospital_id on relations comes from
    the caller (JWT session), not from the doctor record. Also sets Users.hospital_id and
    hospital_name from the session hospital (new users on put_item; existing via update).

    Without doctor_id: create/link a patient to the hospital only via Users.hospital_id /
    hospital_name (from session). No UserDoctorRelations row.
    """
    tail = _phone_tail(phone)
    normalized = normalize_phone_number(phone)
    if not normalized:
        logger.warning(
            f"[staff] add_patient invalid phone hospital_id={hospital_id!r} "
            f"doctor_id={doctor_id!r} raw={phone!r}"
        )
        raise HTTPException(status_code=400, detail=f"Invalid phone number: {phone}")

    if not doctor_id or not str(doctor_id).strip():
        return _add_patient_hospital_only(
            name,
            email,
            hospital_id=hospital_id,
            hospital_name=hospital_name,
            normalized_phone=normalized,
            tail=tail,
            age=age,
            gender=gender,
            insurer=insurer,
        )

    d_id = str(doctor_id).strip()
    doc = (
        preloaded_doctor
        if preloaded_doctor is not None
        else get_doctor_for_hospital_staff_patient_flow(d_id)
    )
    if doc.get("doctor_id") != d_id:
        raise HTTPException(status_code=400, detail="doctor_id does not match preloaded_doctor")
    if doc.get("hospital_id") != hospital_id:
        raise HTTPException(
            status_code=403,
            detail="Doctor does not belong to the hospital in this request",
        )

    logger.info(
        f"[staff] add_patient start doctor_id={d_id!r} hospital_id={hospital_id!r} "
        f"phone={tail} name={name!r}"
    )

    existing = _get_user_by_phone(normalized)
    if existing:
        user_id = existing["user_id"]
        logger.info(
            f"[staff] add_patient existing user doctor_id={d_id!r} phone={tail} "
            f"user_id={user_id!r} ensuring relation and hospital on Users"
        )
        _ensure_auth_record(normalized, "user", user_id=user_id, email=email)
        _ensure_relation(user_id, d_id, hospital_id)
        _ensure_user_hospital_affiliation(user_id, hospital_id, hospital_name)
        merged = {**existing, "hospital_id": hospital_id, "hospital_name": hospital_name}
        return {"status": "linked", "user": merged}

    user_id = generate_user_id()
    now_iso = datetime.now(timezone.utc).isoformat()
    user_item = {
        "user_id": user_id,
        "phoneNumber": normalized,
        "name": name.strip(),
        "source": "hospital_import",
        "createdAt": now_iso,
        "hospital_id": hospital_id,
        "hospital_name": hospital_name,
    }
    if email:
        user_item["email"] = email.lower().strip()
    if age is not None:
        user_item["age"] = age
    if gender is not None and str(gender).strip():
        user_item["gender"] = str(gender).strip()
    if insurer is not None and str(insurer).strip():
        user_item["insurer"] = str(insurer).strip()

    try:
        USERS_TABLE.put_item(Item=user_item)
        logger.info(
            f"[staff] add_patient created user_id={user_id!r} doctor_id={d_id!r} phone={tail}"
        )
    except Exception as e:
        logger.error(
            f"[staff] add_patient Users put_item failed doctor_id={d_id!r} phone={tail} error={e!r}"
        )
        raise HTTPException(status_code=500, detail="Failed to create patient")

    try:
        _ensure_auth_record(normalized, "user", user_id=user_id, email=email)
    except Exception as e:
        logger.error(
            f"[staff] add_patient auth record failed user_id={user_id!r} phone={tail} error={e!r}"
        )
        raise HTTPException(status_code=500, detail="Failed to finalize patient account")

    _ensure_relation(user_id, d_id, hospital_id)

    return {"status": "created", "user": user_item}


_PROTECTED_PATIENT_UPDATE_FIELDS = {
    "user_id",
    "phone",
    "phoneNumber",
    "hospital_id",
    "hospital_name",
    "createdAt",
    "created_at",
    "source",
    "doctor_id",
}


def _clean_patient_update_fields(fields: Dict[str, Any]) -> Dict[str, Any]:
    cleaned: Dict[str, Any] = {}
    for field, value in fields.items():
        if field in _PROTECTED_PATIENT_UPDATE_FIELDS or value is None:
            continue
        if not str(field).strip():
            continue

        if field in {"name", "gender", "insurer"} and isinstance(value, str):
            value = value.strip()
            if not value:
                continue
        if field == "email" and isinstance(value, str):
            value = value.lower().strip()
            if not value:
                continue

        cleaned[field] = value
    return cleaned


def update_patient(
    *,
    hospital_id: str,
    hospital_name: str = "",
    user_id: str,
    doctor_id: Optional[str] = None,
    preloaded_doctor: Optional[dict] = None,
    updates: Optional[Dict[str, Any]] = None,
) -> dict:
    """Partially update an existing hospital patient Users row.

    Doctor assignment rules:
    - doctor_id not provided → doctor relation untouched.
    - doctor_id matches current hospital-assigned doctor → no relation change.
    - doctor_id differs → deactivate old hospital-assigned relation, create new one.
    """
    patient = _get_user_by_id(user_id)
    if not patient:
        raise HTTPException(status_code=404, detail="Patient not found")

    resolved_user_id = patient.get("user_id")
    if not resolved_user_id:
        raise HTTPException(status_code=500, detail="Patient record is missing user_id")

    if patient.get("hospital_id") != hospital_id:
        raise HTTPException(status_code=403, detail="Patient does not belong to your hospital")

    update_fields = _clean_patient_update_fields(updates or {})
    if not update_fields and not (doctor_id and str(doctor_id).strip()):
        raise HTTPException(status_code=400, detail="No patient fields provided to update")

    now_iso = datetime.now(timezone.utc).isoformat()
    update_fields["updatedAt"] = now_iso
    if hospital_name and not patient.get("hospital_name"):
        update_fields["hospital_name"] = hospital_name

    expr_parts = []
    attr_names = {}
    attr_vals = {}
    for idx, (field, value) in enumerate(update_fields.items()):
        name_ph = f"#f{idx}"
        val_ph = f":v{idx}"
        expr_parts.append(f"{name_ph} = {val_ph}")
        attr_names[name_ph] = field
        attr_vals[val_ph] = value

    try:
        resp = USERS_TABLE.update_item(
            Key={"user_id": resolved_user_id},
            UpdateExpression="SET " + ", ".join(expr_parts),
            ExpressionAttributeNames=attr_names,
            ExpressionAttributeValues=attr_vals,
            ReturnValues="ALL_NEW",
        )
        updated_user = resp.get("Attributes") or {**patient, **update_fields}
        logger.info(
            f"[staff] update_patient updated user_id={resolved_user_id!r} "
            f"hospital_id={hospital_id!r} fields={list(update_fields.keys())}"
        )
    except Exception as e:
        logger.error(
            f"[staff] update_patient Users update_item failed user_id={resolved_user_id!r} "
            f"hospital_id={hospital_id!r} error={e!r}"
        )
        raise HTTPException(status_code=500, detail="Failed to update patient")

    # ── Doctor relation swap ────────────────────────────────────────────────
    doctor_action = "none"
    if doctor_id and str(doctor_id).strip():
        d_id = str(doctor_id).strip()
        doc = (
            preloaded_doctor
            if preloaded_doctor is not None
            else get_doctor_for_hospital_staff_patient_flow(d_id)
        )
        if doc.get("doctor_id") != d_id:
            raise HTTPException(status_code=400, detail="doctor_id does not match preloaded_doctor")
        if doc.get("hospital_id") != hospital_id:
            raise HTTPException(status_code=403, detail="Doctor does not belong to your hospital")

        current = get_active_hospital_assigned_doctor(resolved_user_id, hospital_id)
        current_doctor_id = current.get("doctor_id") if current else None

        if current_doctor_id == d_id:
            doctor_action = "unchanged"
            logger.info(
                f"[staff] update_patient doctor unchanged user_id={resolved_user_id!r} "
                f"doctor_id={d_id!r}"
            )
        else:
            if current_doctor_id:
                try:
                    deactivate_relation(resolved_user_id, current_doctor_id)
                    logger.info(
                        f"[staff] update_patient deactivated old relation "
                        f"user_id={resolved_user_id!r} old_doctor_id={current_doctor_id!r}"
                    )
                except Exception as e:
                    logger.warning(
                        f"[staff] update_patient could not deactivate old relation "
                        f"user_id={resolved_user_id!r} old_doctor_id={current_doctor_id!r} err={e!r}"
                    )
            _ensure_relation(resolved_user_id, d_id, hospital_id)
            doctor_action = "updated"
            logger.info(
                f"[staff] update_patient doctor updated user_id={resolved_user_id!r} "
                f"new_doctor_id={d_id!r} old_doctor_id={current_doctor_id!r}"
            )

    return {
        "status": "updated",
        "user": updated_user,
        "updated_fields": sorted(update_fields.keys()),
        "doctor_action": doctor_action,
    }


def _add_patient_hospital_only(
    name: str,
    email: Optional[str],
    *,
    hospital_id: str,
    hospital_name: str,
    normalized_phone: str,
    tail: str,
    age: Optional[int] = None,
    gender: Optional[str] = None,
    insurer: Optional[str] = None,
) -> dict:
    """Patient under hospital only: Users.hospital_id / hospital_name; no user–doctor relation."""
    logger.info(
        f"[staff] add_patient hospital-only hospital_id={hospital_id!r} phone={tail} name={name!r}"
    )

    existing = _get_user_by_phone(normalized_phone)
    if existing:
        user_id = existing["user_id"]
        _ensure_auth_record(normalized_phone, "user", user_id=user_id, email=email)
        _ensure_user_hospital_affiliation(user_id, hospital_id, hospital_name)
        merged = {**existing, "hospital_id": hospital_id, "hospital_name": hospital_name}
        return {"status": "linked", "user": merged}

    user_id = generate_user_id()
    now_iso = datetime.now(timezone.utc).isoformat()
    user_item = {
        "user_id": user_id,
        "phoneNumber": normalized_phone,
        "name": name.strip(),
        "source": "hospital_import",
        "createdAt": now_iso,
        "hospital_id": hospital_id,
        "hospital_name": hospital_name,
    }
    if email:
        user_item["email"] = email.lower().strip()
    if age is not None:
        user_item["age"] = age
    if gender is not None and str(gender).strip():
        user_item["gender"] = str(gender).strip()
    if insurer is not None and str(insurer).strip():
        user_item["insurer"] = str(insurer).strip()

    try:
        USERS_TABLE.put_item(Item=user_item)
        logger.info(
            f"[staff] add_patient hospital-only created user_id={user_id!r} phone={tail}"
        )
    except Exception as e:
        logger.error(
            f"[staff] add_patient hospital-only put_item failed phone={tail} error={e!r}"
        )
        raise HTTPException(status_code=500, detail="Failed to create patient")

    try:
        _ensure_auth_record(normalized_phone, "user", user_id=user_id, email=email)
    except Exception as e:
        logger.error(
            f"[staff] add_patient hospital-only auth record failed user_id={user_id!r} error={e!r}"
        )
        raise HTTPException(status_code=500, detail="Failed to finalize patient account")

    return {"status": "created", "user": user_item}


def _ensure_relation(user_id: str, doctor_id: str, hospital_id: str) -> None:
    """Best-effort: HOSPITAL_ASSIGNED relation; hospital_id from staff session (UserDoctorRelations)."""
    try:
        create_relation(
            user_id=user_id,
            doctor_id=doctor_id,
            relation_type=RelationType.HOSPITAL_ASSIGNED,
            linked_by=LinkedBy.HOSPITAL_STAFF,
            hospital_id=hospital_id,
        )
    except Exception as e:
        logger.error(
            f"[staff] _ensure_relation failed user_id={user_id!r} doctor_id={doctor_id!r} "
            f"hospital_id={hospital_id!r} error={e!r}"
        )


# ---------------------------------------------------------------------------
# Core: add doctor
# ---------------------------------------------------------------------------

def add_doctor(
    hospital_id: str,
    hospital_name: str,
    phone: str,
    name: str,
    specialization: str,
    experience: str,
    email: Optional[str] = None,
) -> dict:
    tail = _phone_tail(phone)
    logger.info(
        f"[staff] add_doctor start hospital_id={hospital_id!r} phone={tail} name={name!r} "
        f"specialization={specialization!r}"
    )
    normalized = normalize_phone_number(phone)
    if not normalized:
        logger.warning(f"[staff] add_doctor invalid phone hospital_id={hospital_id!r} raw={phone!r}")
        raise HTTPException(status_code=400, detail=f"Invalid phone number: {phone}")

    existing = _get_doctor_by_phone(normalized)
    if existing:
        doctor_id = existing["doctor_id"]
        logger.info(
            f"[staff] add_doctor existing doctor hospital_id={hospital_id!r} phone={tail} "
            f"doctor_id={doctor_id!r} linking"
        )
        _link_doctor_to_hospital(doctor_id, hospital_id, hospital_name)
        _ensure_auth_record(normalized, "doctor", doctor_id=doctor_id, email=email)
        return {"status": "linked", "doctor": existing}

    doctor_id = generate_doctor_id()
    now_iso = datetime.now(timezone.utc).isoformat()
    doctor_item = {
        "doctor_id": doctor_id,
        "phoneNumber": normalized,
        "doctorname": name.strip(),
        "specialization": specialization,
        "experience": experience,
        "hospital_id": hospital_id,
        "hospital_name": hospital_name,
        "source": "hospital_import",
        "createdAt": now_iso,
    }
    if email:
        doctor_item["email"] = email.lower().strip()

    try:
        DOCTORS_TABLE.put_item(Item=doctor_item)
        logger.info(
            f"[staff] add_doctor created doctor_id={doctor_id!r} hospital_id={hospital_id!r} phone={tail}"
        )
    except Exception as e:
        logger.error(
            f"[staff] add_doctor Doctors put_item failed hospital_id={hospital_id!r} phone={tail} error={e!r}"
        )
        raise HTTPException(status_code=500, detail="Failed to create doctor")

    try:
        _ensure_auth_record(normalized, "doctor", doctor_id=doctor_id, email=email)
    except Exception as e:
        logger.error(
            f"[staff] add_doctor auth record failed doctor_id={doctor_id!r} phone={tail} error={e!r}"
        )
        raise HTTPException(status_code=500, detail="Failed to finalize doctor account")
    _create_default_slots(doctor_id)
    return {"status": "created", "doctor": doctor_item}


def _link_doctor_to_hospital(doctor_id: str, hospital_id: str, hospital_name: str) -> None:
    try:
        DOCTORS_TABLE.update_item(
            Key={"doctor_id": doctor_id},
            UpdateExpression="SET hospital_id = :hid, hospital_name = :hname",
            ExpressionAttributeValues={
                ":hid": hospital_id,
                ":hname": hospital_name,
            },
        )
        logger.info(
            f"[staff] linked doctor doctor_id={doctor_id!r} to hospital_id={hospital_id!r} "
            f"hospital_name={hospital_name!r}"
        )
    except Exception as e:
        logger.error(
            f"[staff] link doctor failed doctor_id={doctor_id!r} hospital_id={hospital_id!r} error={e!r}"
        )


# ---------------------------------------------------------------------------
# Bulk: import patients from Excel (openpyxl only — avoids pandas/numpy in Lambda)
# ---------------------------------------------------------------------------

def _xlsx_to_records(content: bytes) -> tuple[list[str], list[dict]]:
    """First row = headers (lowercased); following rows = dicts. .xlsx only."""
    from openpyxl import load_workbook

    size_kb = len(content) / 1024.0
    logger.info(f"[staff] _xlsx_to_records parsing xlsx size_kb={size_kb:.2f}")
    bio = io.BytesIO(content)
    wb = load_workbook(bio, read_only=True, data_only=True)
    try:
        ws = wb.active
        it = ws.iter_rows(values_only=True)
        first = next(it, None)
        if first is None:
            logger.warning("[staff] _xlsx_to_records empty workbook (no header row)")
            raise HTTPException(status_code=400, detail="Excel file is empty")
        headers = []
        for c in first:
            headers.append(str(c).strip().lower() if c is not None else "")
        records = []
        for row in it:
            rec = {}
            for i, h in enumerate(headers):
                if not h:
                    continue
                rec[h] = row[i] if i < len(row) else None
            records.append(rec)
        logger.info(
            f"[staff] _xlsx_to_records parsed headers={headers!r} data_rows={len(records)}"
        )
        return headers, records
    finally:
        wb.close()


def import_patients_from_excel(
    preloaded_doctor: dict,
    file: UploadFile,
) -> dict:
    fname = getattr(file, "filename", None) or ""
    doctor_id = preloaded_doctor["doctor_id"]
    hid = preloaded_doctor.get("hospital_id", "")
    logger.info(
        f"[staff] import_patients_from_excel start doctor_id={doctor_id!r} hospital_id={hid!r} filename={fname!r}"
    )
    if not file.filename or not file.filename.lower().endswith(".xlsx"):
        logger.warning(f"[staff] import_patients wrong extension filename={fname!r}")
        raise HTTPException(status_code=400, detail="File must be .xlsx")

    try:
        content = file.file.read()
        if not content:
            logger.warning("[staff] import_patients empty file body")
            raise HTTPException(status_code=400, detail="Uploaded file is empty")
    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"[staff] import_patients read failed filename={fname!r} error={e!r}")
        raise HTTPException(status_code=400, detail=f"Failed to read file: {e}")

    try:
        headers, rows = _xlsx_to_records(content)
    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"[staff] import_patients parse failed filename={fname!r} error={e!r}")
        raise HTTPException(status_code=400, detail=f"Invalid Excel file: {e}")

    if "phone" not in headers:
        logger.warning(f"[staff] import_patients missing phone column headers={headers!r}")
        raise HTTPException(status_code=400, detail="Excel must contain a 'phone' column")
    if "name" not in headers:
        logger.warning(f"[staff] import_patients missing name column headers={headers!r}")
        raise HTTPException(status_code=400, detail="Excel must contain a 'name' column")

    summary = {"total_rows": len(rows), "created": 0, "linked": 0, "skipped": 0, "errors": []}

    for idx, row in enumerate(rows, start=2):
        phone_raw = row.get("phone")
        name_raw = row.get("name")
        email_raw = row.get("email")

        if _is_nan(phone_raw) or not str(phone_raw).strip():
            summary["skipped"] += 1
            continue
        if _is_nan(name_raw) or not str(name_raw).strip():
            summary["skipped"] += 1
            continue

        email = str(email_raw).strip() if not _is_nan(email_raw) else None

        try:
            result = add_patient(
                str(phone_raw).strip(),
                str(name_raw).strip(),
                email,
                hospital_id=hid,
                hospital_name=preloaded_doctor.get("hospital_name") or "",
                doctor_id=doctor_id,
                preloaded_doctor=preloaded_doctor,
            )
            if result["status"] == "created":
                summary["created"] += 1
            else:
                summary["linked"] += 1
        except HTTPException as e:
            summary["errors"].append({"row": idx, "detail": e.detail})
            logger.warning(
                f"[staff] import_patients row error doctor_id={doctor_id!r} row={idx} "
                f"HTTP {e.status_code}: {e.detail}"
            )
        except Exception as e:
            summary["errors"].append({"row": idx, "detail": str(e)})
            logger.warning(
                f"[staff] import_patients row error doctor_id={doctor_id!r} row={idx} error={e!r}"
            )

    err_n = len(summary["errors"])
    if err_n:
        logger.warning(
            f"[staff] import_patients finished with row errors doctor_id={doctor_id!r} "
            f"errors={err_n} sample={summary['errors'][:5]}"
        )
    logger.info(
        f"[staff] import_patients summary doctor_id={doctor_id!r} "
        f"total_rows={summary['total_rows']} created={summary['created']} "
        f"linked={summary['linked']} skipped={summary['skipped']} errors={err_n}"
    )
    return summary


# ---------------------------------------------------------------------------
# Bulk: import doctors from Excel
# ---------------------------------------------------------------------------

def import_doctors_from_excel(
    hospital_id: str,
    hospital_name: str,
    file: UploadFile,
) -> dict:
    fname = getattr(file, "filename", None) or ""
    logger.info(
        f"[staff] import_doctors_from_excel start hospital_id={hospital_id!r} filename={fname!r}"
    )
    if not file.filename or not file.filename.lower().endswith(".xlsx"):
        logger.warning(f"[staff] import_doctors wrong extension filename={fname!r}")
        raise HTTPException(status_code=400, detail="File must be .xlsx")

    try:
        content = file.file.read()
        if not content:
            logger.warning("[staff] import_doctors empty file body")
            raise HTTPException(status_code=400, detail="Uploaded file is empty")
    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"[staff] import_doctors read failed filename={fname!r} error={e!r}")
        raise HTTPException(status_code=400, detail=f"Failed to read file: {e}")

    try:
        headers, rows = _xlsx_to_records(content)
    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"[staff] import_doctors parse failed filename={fname!r} error={e!r}")
        raise HTTPException(status_code=400, detail=f"Invalid Excel file: {e}")

    required = ["phone", "name", "specialization", "experience"]
    missing = [c for c in required if c not in headers]
    if missing:
        logger.warning(f"[staff] import_doctors missing columns {missing} headers={headers!r}")
        raise HTTPException(
            status_code=400,
            detail=f"Excel is missing required columns: {', '.join(missing)}",
        )

    summary = {"total_rows": len(rows), "created": 0, "linked": 0, "skipped": 0, "errors": []}

    for idx, row in enumerate(rows, start=2):
        phone_raw = row.get("phone")
        name_raw = row.get("name")
        spec_raw = row.get("specialization")
        exp_raw = row.get("experience")
        email_raw = row.get("email")

        if _is_nan(phone_raw) or not str(phone_raw).strip():
            summary["skipped"] += 1
            continue
        if _is_nan(name_raw) or not str(name_raw).strip():
            summary["skipped"] += 1
            continue
        if _is_nan(spec_raw) or not str(spec_raw).strip():
            summary["skipped"] += 1
            continue
        if _is_nan(exp_raw):
            summary["skipped"] += 1
            continue

        email = str(email_raw).strip() if not _is_nan(email_raw) else None

        try:
            result = add_doctor(
                hospital_id=hospital_id,
                hospital_name=hospital_name,
                phone=str(phone_raw).strip(),
                name=str(name_raw).strip(),
                specialization=str(spec_raw).strip(),
                experience=str(exp_raw).strip(),
                email=email,
            )
            if result["status"] == "created":
                summary["created"] += 1
            else:
                summary["linked"] += 1
        except HTTPException as e:
            summary["errors"].append({"row": idx, "detail": e.detail})
            logger.warning(
                f"[staff] import_doctors row error hospital_id={hospital_id!r} row={idx} "
                f"HTTP {e.status_code}: {e.detail}"
            )
        except Exception as e:
            summary["errors"].append({"row": idx, "detail": str(e)})
            logger.warning(
                f"[staff] import_doctors row error hospital_id={hospital_id!r} row={idx} error={e!r}"
            )

    err_n = len(summary["errors"])
    if err_n:
        logger.warning(
            f"[staff] import_doctors finished with row errors hospital_id={hospital_id!r} "
            f"errors={err_n} sample={summary['errors'][:5]}"
        )
    logger.info(
        f"[staff] import_doctors summary hospital_id={hospital_id!r} "
        f"total_rows={summary['total_rows']} created={summary['created']} "
        f"linked={summary['linked']} skipped={summary['skipped']} errors={err_n}"
    )
    return summary


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _is_nan(value) -> bool:
    """Check if a value is NaN or None."""
    if value is None:
        return True
    if isinstance(value, float):
        import math
        return math.isnan(value)
    return False
